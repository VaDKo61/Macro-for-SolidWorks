import os

import openpyxl
import win32com.client
from openpyxl.styles import Font, Alignment

from functions.archive.create_drill_sheet import get_ready


def create_specification(sw_app, sw_model):
    """create drawing and save specification in excel"""
    assembly_path: str = sw_model.GetPathName
    assembly_path_list: list = assembly_path.split('\\')[2:]
    assembly_name: str = assembly_path_list[-1].split('.')[0]
    engineer: str = assembly_path_list[3]

    # select frame part
    frames: list = []
    frame_names: list = []
    for component in sw_model.GetComponents(True):
        if 'Рама' in component.Name2:
            frames.append(component.GetPathName)
            frame_names.append('-'.join(component.Name2.split('-')[0:-1]))

    # create draw
    template_path = f'\\\\{assembly_path_list[0]}\\{assembly_path_list[1]}\\{assembly_path_list[2]}\\{engineer}\\' \
                    f'Шаблоны\\Чертеж спецификации.DRWDOT'
    sw_draw = sw_app.NewDocument(template_path, 12, 0.42, 0.297)

    # add draw view
    x_view: float = 0
    y_view: float = 0.148
    frames_view: list = []
    assembly_view = sw_draw.CreateDrawViewFromModelView3(assembly_path, '*Изометрия', x_view, y_view, 0)
    for frame in frames:
        x_view += 0.52
        frames_view.append(sw_draw.CreateDrawViewFromModelView3(frame, '*Изометрия', x_view, y_view, 0))

    # add specification
    template_assembly = '\\\\192.168.1.14\\SolidWorks\\Библиотека Solid Works НОВАЯ\\Шаблоны\\' \
                        'Шаблон специи.sldbomtbt'
    template_frame = '\\\\192.168.1.14\\SolidWorks\\Библиотека Solid Works НОВАЯ\\Шаблоны\\' \
                     'Сварные конструкции.sldwldtbt'
    x_spec: float = -0.2
    y_spec: float = -0.03
    if get_ready():
        return
    table_assembly = assembly_view.InsertBomTable4(False, x_spec, y_spec, 1, 1, 'По умолчанию', template_assembly,
                                                   False, 0, False)
    tables_frame: list = []
    for view in frames_view:
        x_spec += 0.52
        tables_frame.append(
            view.InsertWeldmentTable(False, x_spec, y_spec, 1, 'По умолчанию<Как сварной>', template_frame))

    # save draw
    path_specification: str = '\\'.join(assembly_path.split('\\')[:-1]) + '\\Спецификации'
    if not os.path.isdir(path_specification):
        os.makedirs(path_specification)
    path_draw = f'{path_specification}\\{assembly_name} (Cпецификация).SLDDRW'
    sw_draw.SaveAs3(path_draw, 0, 0)

    # save table excel
    table_assembly.SaveAsExcel(f'{path_specification}//{assembly_name}.xlsx', False, False)
    for i, table in enumerate(tables_frame, start=1):
        text = table.WeldmentCutListFeature.GetTableAnnotations[0]
        text.SaveAsText2(f'{path_specification}//{frame_names[i - 1]}.txt', '__', False)
    count = 0
    for file in os.listdir(f'{path_specification}'):
        if file.split('.')[-1] == 'txt':
            count += 1
            with open(f'{path_specification}\\{file}', 'r') as file_text:
                content: str = file_text.read()
                content_list: list = [i.split('__') for i in content.split('\n') if i]
                for row in content_list:
                    for i, cell in enumerate(row):
                        if cell.isdigit():
                            row[i] = int(cell)
                        if ''.join(cell.split('.')).isdigit():
                            row[i] = float(cell)
                excel_list: list[tuple] = list(map(tuple, content_list))
            wb = openpyxl.Workbook()
            ws = wb.active
            for row in excel_list:
                ws.append(row)
            name_size_column = {
                'A': 5,
                'B': 16.5,
                'C': 80,
                'D': 13.5,
                'E': 10.4,
                'F': 11.8,
                'G': 25,
                'H': 22.5,
                'I': 21.3,
                'J': 11
            }
            for name_column, size in name_size_column.items():
                ws.column_dimensions[f'{name_column}'].width = size
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = Font(name='Times New Roman', size=14, vertAlign='baseline')
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            wb.save(f'{path_specification}//{"".join(file.split(".")[0:-1])}.xlsx')
            os.remove(f'{path_specification}\\{file}')


def specification():
    sw_app = win32com.client.dynamic.Dispatch('SldWorks.Application')
    sw_model = sw_app.ActiveDoc
    if sw_model.GetType != 2:
        sw_app.SendmsgToUser('Активна не сборка')
        return
    create_specification(sw_app, sw_model)
    sw_app.SendmsgToUser('Чертеж создан и сохранены Excel')
